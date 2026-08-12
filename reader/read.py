import bs4
import docx
import fitz
import html
import json
import transliterate
import re


def close_index(text, key_word):
    index_pos = text.find(key_word)
    global_end = 0
    word_length = len(key_word)

    while index_pos != -1:
        global_index_pos = global_end + index_pos
        global_end = global_index_pos + text[global_index_pos:].find('<')
        text = text[:global_end] + key_word + text[global_end:]
        index_pos = text[global_end + word_length:].find(key_word)

    return text

def convert_pdf_to_html(pdf_path, output_path):
    file = fitz.open(pdf_path)

    html_parts = ["<html><body>"]
    text = html.unescape(file[0].get_text("html"))

    text = text.replace(f'font-size:{title_sub_font}pt;color:#231f20">',
                        f'font-size:{title_font}pt;color:#231f20">down_index')
    text = text.replace(f'font-size:{annot_sub_font}pt;color:#231f20">',
                        f'font-size:{annot_font}pt;color:#231f20">down_index')
    text = text.replace(f'font-size:{annot_sub_2_font}pt;color:#231f20">',
                        f'font-size:{annot_font}pt;color:#231f20">down_index')

    text = text.replace(
        f'<sup><i><span style="font-family:Times New Roman,serif;font-size:{title_font}pt;color:#231f20">down_index',
        f'<sup><i><span style="font-family:Times New Roman,serif;font-size:{title_font}pt;color:#231f20">up_index')
    text = text.replace(
        f'<sup><b><span style="font-family:Times New Roman,serif;font-size:{title_font}pt;color:#231f20">down_index',
        f'<sup><b><span style="font-family:Times New Roman,serif;font-size:{title_font}pt;color:#231f20">up_index')
    text = text.replace(
        f'<sup><i><span style="font-family:Times New Roman,serif;font-size:{annot_font}pt;color:#231f20">down_index',
        f'<sup><i><span style="font-family:Times New Roman,serif;font-size:{annot_font}pt;color:#231f20">up_index')
    text = text.replace(
        f'<sup><b><span style="font-family:Times New Roman,serif;font-size:{annot_font}pt;color:#231f20">down_index',
        f'<sup><b><span style="font-family:Times New Roman,serif;font-size:{annot_font}pt;color:#231f20">up_index')

    text = text.replace('´', 'x')

    html_parts.append(text)
    html_parts.append("</body></html>")
    text = close_index(close_index('\n'.join(html_parts), 'up_index'), 'down_index')

    with open(output_path, "w", encoding="utf-8") as f:
        f.write(text)

    return len(file)

def extract(soup, font):
    result = ''

    for element in soup.select('[style*=font-size]'):
        if float(element['style'].split(';')[1][10: -2]) == font:
            result += element.text.replace('\n', ' ').replace('  ', ' ')

    return result

def dump_subscript(par, part, index, key_word):
    par.add_run(part[:index])
    next = index + 1
    end = next + part[next:].find(key_word)
    length = len(key_word)

    if key_word == 'down_index':
        par.add_run(part[index + length: end]).font.subscript = True
    else:
        par.add_run(part[index + length: end]).font.superscript = True

    part = part[end + length:]
    return part

def insert_with_subscripts(part):
    par = info.add_paragraph()


    while True:
        down_index = part.find('down_index')
        up_index = part.find('up_index')
        index = 0

        if down_index != -1 and up_index != -1:
            index = min(down_index,  up_index)
        elif down_index == -1 and up_index != -1:
            index = up_index
        elif down_index != -1 and up_index == -1:
            index = down_index
        else:
            break

        if index == down_index:
            part = dump_subscript(par, part, index, 'down_index')
        else:
            part = dump_subscript(par, part, index, 'up_index')

    par.add_run(part)


path = 'C:/Users/gubar/OneDrive/Документы/ФИАН/Квантовая Электроника'
pdf_path = 'pdf/to_convert'
out_folder = 'html'
docx_folder = 'docx/9'
name = '533'

title_font = 17
title_sub_font = 11.9
author_font = 11.5
annot_font = 8.7
annot_sub_font = 6.1
annot_sub_2_font = 4.8
affiliation_font = 8
numbers_font = 7.5

need_dates = False
need_annot = True
need_save = True

out_path = f'{path}/{out_folder}/{name}.html'
pages = convert_pdf_to_html(f'{path}/{pdf_path}/{name}.pdf', out_path)

with open(out_path, encoding="utf-8") as html_file:
    soup = bs4.BeautifulSoup(html_file, 'html.parser')

title = extract(soup, title_font)

if title[-1] == ' ':
    title = title[:-1]

authors = extract(soup, author_font).replace('  ', ' ').replace('. ', '.')

if authors[-1] == ' ':
    authors = authors[:-1]

authors_with_seps = authors.replace('.', '. ')

all_affiliation_records = []
all_affiliations = set()
affiliation = ''

for element in soup.select('[style*=font-size]'):
    if float(element['style'].split(';')[1][10: -2]) == affiliation_font:
        text = element.text

        if text == ' ':
            all_affiliation_records.append(affiliation)
            affiliation = ''
        else:
            affiliation += text

if affiliation.find('Поступила') != -1:
    affiliations_and_dates = affiliation.split('Поступила')
else:
    affiliations_and_dates = affiliation.split('Поступило')

all_affiliation_records.append(affiliations_and_dates[0])

for element in all_affiliation_records:
    element = element.replace('­', '').replace('- \n', '')
    for author in authors.split(', '):
        element = element.replace(f'{author}, ', '\t').replace(f'{author}. ', '\t').replace(' ', ' ')

    for place in re.split(r'[;\t]', element):
        if len(place) > 0 and place[0] == ' ':
            place = place[1:]
        if len(place) > 0 and place[0] == '	':
            place = place[1:]
        if len(place) > 0 and place[-1] == ' ':
            place = place[:-1]

        if len(place) > 0:
            if 'e-mail' not in place and '@' not in place:
                all_affiliations.add(place.replace('c', 'с').replace('-', ''))

if need_annot:
    annot_and_key = extract(soup, annot_font).replace('­', '').split('.')
    annot = '.'.join(annot_and_key[:-2]) + '.'
    key = annot_and_key[-2]

    if key[0] == ' ':
        key = key[1:]

    print(annot)
    print(key)

to_translit = title.replace('.', '').replace(',', '').replace('/', '')
to_translit = to_translit.replace(':', '').replace('%', '').replace(' ', '')
to_translit = to_translit.replace('+', '')

translit = transliterate.translit(to_translit, 'ru',
                                  reversed=True).replace(' ', '-').replace("'", '').lower()
translit = translit.replace('down_index', '').replace('up_index', '')

print(title)
print(translit)
print(authors_with_seps)

for place in all_affiliations:
    print(place)

if need_dates:
    dates = affiliations_and_dates[1].replace('­', '').replace(' ', ' ').replace('  ', ' ').split(',')[:3]
    month = {'января': '01', 'февраля': '02', 'марта': '03', 'апреля': '04', 'мая': '05', 'июня': '06', 'июля': '07',
             'августа': '08', 'сентября': '09', 'октября': '10', 'ноября': '11', 'декабря': '12'}

    for index, date_record in enumerate(dates):
        date_record = date_record[:date_record.find('.')]
        date = date_record.split(' ')[-4: -1]
        date[1] = month[date[1]]
        dates[index] = '.'.join(date)

    for date in dates:
        print(date)

cite = f'{authors_with_seps}, "{title}"'
print(cite)
issue_year_and_page = extract(soup, numbers_font).split('ke@lebedev.ru')

if issue_year_and_page[1].find('«') != -1:
    issue_year_and_page.reverse()

issue_and_year = issue_year_and_page[0].split(', ')
extract_start_page = issue_year_and_page[1]
start_page = int(extract_start_page[extract_start_page.find('\t') + 1:])

issue = issue_and_year[1]
year = issue_and_year[2].replace('№ ', '').replace('	', '')
pages_record = f'{start_page}-{start_page + pages - 1}'

print(issue)
print(year)
print(pages_record)

if need_save:
    info = docx.Document()

    insert_with_subscripts(title)
    info.add_paragraph(translit)
    info.add_paragraph(authors_with_seps)
    info.add_paragraph('<ul>')

    for place in all_affiliations:
        info.add_paragraph(f'<li> {place} </li>')

    info.add_paragraph('</ul>')

    if need_annot:
        insert_with_subscripts(f'<div class = "around-button"> <b> Аннотация: </b> {annot} </div>')
        insert_with_subscripts(f'<div class = "around-button"> <b> Ключевые слова: </b> <i> {key} </i> </div>')

    info.add_paragraph('<div> </div>')
    info.add_paragraph('<div>')

    if need_dates:
        info.add_paragraph(f'<div class = "around-button"> <b> Поступила в редакцию: </b> {dates[0]} </div>')

        if len(dates) == 3:
            info.add_paragraph(f'<div class = "around-button"> <b> После доработки: </b> {dates[1]} </div>')

        if len(dates) > 1:
            info.add_paragraph(f'<div class = "around-button"> <b> Принята в печать: </b> {dates[-1]} </div>')
    else:
        info.add_paragraph('<div class = "around-button"> <b> Поступила в редакцию: </b>  </div>')
        info.add_paragraph('<div class = "around-button"> <b> После доработки: </b>  </div>')
        info.add_paragraph('<div class = "around-button"> <b> Принята в печать: </b>  </div>')

    info.add_paragraph('</div>')
    info.add_paragraph('<div> </div>')
    insert_with_subscripts(f'<div> <b> Образец цитирования: </b> {cite}, <i> Квантовая электроника </i>, '
                       f'<b> {issue} </b>: {year}, {pages_record} </div>')
    info.add_paragraph('Скачать (.pdf)')

    info.save(f'{path}/{docx_folder}/{start_page}.docx')
